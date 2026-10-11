"""Run each project's code and record what it actually prints or returns.

The home page shows, for every project, a code excerpt next to the REAL output
of that code. This script produces those outputs, so they are never typed by hand:

  * SQL     - runs the risk-scoring query on the journal entry database (SQLite);
  * Excel   - reads the reconciliation statement as last calculated by Excel;
  * Power BI- the DAX query result returned by the Power BI engine during the model test;
  * Python  - runs the fraud pipeline and keeps its console output;
  * R       - runs the housing analysis script and keeps its console output.

Output: assets/data/code_outputs.js (window.CODE_OUTPUTS = {...})

Usage (from the portfolio folder):
    python tools/capture_code_outputs.py
"""

from __future__ import annotations

import json
import os
import sqlite3
import subprocess
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PROJECTS = ROOT / "projects"
RSCRIPT = Path(os.environ.get("RSCRIPT", r"C:\Users\w-Mau\AppData\Local\Programs\R\R-4.6.1\bin\Rscript.exe"))


def box_table(columns, rows, align_right=()):
    """Format rows like the sqlite3 command-line tool in '.mode box'."""
    text_rows = [["" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v)) for v in row] for row in rows]
    widths = [max(len(str(c)), *(len(r[i]) for r in text_rows)) for i, c in enumerate(columns)]

    def line(left, mid, right):
        return left + mid.join("─" * (w + 2) for w in widths) + right

    def cells(values):
        out = []
        for i, (value, width) in enumerate(zip(values, widths)):
            out.append(" " + (value.rjust(width) if i in align_right else value.ljust(width)) + " ")
        return "│" + "│".join(out) + "│"

    return "\n".join([line("┌", "┬", "┐"), cells([str(c) for c in columns]), line("├", "┼", "┤"),
                      *[cells(r) for r in text_rows], line("└", "┴", "┘")])


def run(command, cwd):
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    result = subprocess.run(command, cwd=cwd, capture_output=True, text=True, encoding="utf-8", errors="replace", env=env)
    if result.returncode != 0:
        raise SystemExit(f"{command} failed:\n{result.stderr[-2000:]}")
    return result.stdout


# --- SQL -------------------------------------------------------------------------------

SQL_QUERY = """SELECT
    h.entry_number,
    e.created_by,
    e.total_debit                                    AS amount,
    SUM(c.weight)                                    AS risk_score,
    GROUP_CONCAT(h.test_id, ', ' ORDER BY h.test_id) AS tests
FROM test_hits AS h
JOIN test_catalog AS c ON c.test_id = h.test_id
JOIN v_entries    AS e ON e.entry_number = h.entry_number
GROUP BY h.entry_number, e.created_by, e.total_debit
ORDER BY risk_score DESC, amount DESC
LIMIT 6;"""


def capture_sql():
    project = PROJECTS / "journal-entry-testing"
    run([sys.executable, "src/run_analysis.py"], project)  # rebuilds output/ledger.db and test_hits
    conn = sqlite3.connect(project / "output" / "ledger.db")
    cursor = conn.execute(SQL_QUERY)
    columns = [d[0] for d in cursor.description]
    rows = cursor.fetchall()
    conn.close()
    return {
        "language": "sql",
        "label": "SQL · SQLite",
        "code": SQL_QUERY,
        "output": box_table(columns, rows, align_right=(2, 3)),
        "source": "Result of the query above on output/ledger.db, after python src/run_analysis.py",
    }


# --- Excel ----------------------------------------------------------------------------------

EXCEL_CODE = """// Power Query (M): read the bank's French CSV export
Typed = Table.TransformColumnTypes(BlanksToNull, {{"Date", type date},
          {"Debit", type number}, {"Credit", type number}}, "fr-FR"),
WithReference = Table.AddColumn(WithAmount, "Reference", each
          if Text.Contains([Label], "/REF ")
          then Text.Trim(Text.AfterDelimiter([Label], "/REF ")) else "")

' Excel: match on amount and date when the bank lost the reference
=IFERROR(INDEX(tblLedger[Entry], MATCH(1, INDEX((tblLedger[Amount]=[@Amount])
   *(tblLedger[MatchedByRef]=FALSE)*(ABS(tblLedger[Date]-[@Date])<=DateTolerance),0),0)),"")"""


def capture_excel():
    from openpyxl import load_workbook

    workbook = load_workbook(PROJECTS / "bank-reconciliation" / "bank_reconciliation.xlsx", data_only=True)
    sheet = workbook["Reconciliation"]
    rows = []
    for r in range(6, 19):
        label, amount, items = (sheet.cell(r, c).value for c in (2, 3, 4))
        if label is None and amount is None:
            continue
        rows.append([label or "", "" if amount is None else amount, "" if items in (None, "") else items])
    lines = []
    for label, amount, items in rows:
        if amount == "":
            lines.append(label)
            continue
        value = f"{amount:>14,.2f}" if isinstance(amount, (int, float)) else str(amount)
        extra = f"   {int(items)} item(s)" if isinstance(items, (int, float)) else (f"   {items}" if items else "")
        lines.append(f"  {label:<66}{value}{extra}")
    return {
        "language": "excel",
        "label": "Excel · Power Query",
        "code": EXCEL_CODE,
        "output": "\n".join(lines),
        "source": "Reconciliation sheet as last calculated by Excel (values read from bank_reconciliation.xlsx)",
    }


# --- Power BI ---------------------------------------------------------------------------------

DAX_QUERY = """EVALUATE
CALCULATETABLE(
    ROW(
        "Net Sales", [Net Sales], "Net Sales LY", [Net Sales LY],
        "Net Sales YoY %", [Net Sales YoY %], "Gross Margin %", [Gross Margin %],
        "Gross Margin % LY", [Gross Margin % LY], "Budget", [Budget],
        "Net Sales vs Budget %", [Net Sales vs Budget %],
        "Gross Margin vs Budget", [Gross Margin vs Budget]
    ),
    dim_date[year] = 2025
)"""


def capture_powerbi():
    check = json.loads((PROJECTS / "sales-dashboard-powerbi" / "results" / "dax_check.json").read_text(encoding="utf-8-sig"))
    totals = check["totals"]
    rows = []
    for name in ["Net Sales", "Net Sales LY", "Net Sales YoY %", "Gross Margin %", "Gross Margin % LY",
                 "Budget", "Net Sales vs Budget %", "Gross Margin vs Budget"]:
        value = totals[f"[{name}]"]
        rows.append([f"[{name}]", f"{value * 100:,.2f} %" if "%" in name else f"{value:,.2f}"])
    return {
        "language": "dax",
        "label": "Power BI · DAX",
        "code": DAX_QUERY,
        "output": box_table(["Measure", "Value"], rows, align_right=(1,)),
        "source": "Returned by the Power BI engine (Analysis Services) when the model was deployed and refreshed",
    }


# --- Python -----------------------------------------------------------------------------------

PYTHON_CODE = """for name, model in make_models().items():
    model.fit(train[features], train["is_fraud"])
    # threshold chosen on April only, on business cost
    threshold, _ = best_threshold(model.predict_proba(validation[features])[:, 1], validation)
    scores = model.predict_proba(test[features])[:, 1]
    results[name] = outcome(scores >= threshold, test, test_days)"""


def capture_python():
    output = run([sys.executable, "src/fraud_pipeline.py"], PROJECTS / "fraud-detection-python")
    lines = [l for l in output.splitlines() if not l.startswith("Recall by pattern")]
    return {
        "language": "python",
        "label": "Python · scikit-learn",
        "code": PYTHON_CODE,
        "output": "$ python src/fraud_pipeline.py\n" + "\n".join(lines),
        "source": "Console output of python src/fraud_pipeline.py (test months May-June 2025)",
    }


# --- R -------------------------------------------------------------------------------------------

R_CODE = """index_model <- lm(log(price_m2) ~ factor(year) + type + log(surface) + rooms + commune,
                  data = sales)
year_effects <- tidy(index_model) |>
  filter(grepl("factor\\\\(year\\\\)", term)) |>
  transmute(year = as.integer(sub("factor\\\\(year\\\\)", "", term)),
            hedonic = exp(estimate) * 100)       # 2021 = 100"""


def capture_r():
    output = run([str(RSCRIPT), "R/02_analysis.R"], PROJECTS / "lille-housing-prices")
    lines = output.splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith("Index (2021 = 100)"))
    end = next(i for i, l in enumerate(lines) if l.startswith("Most / least expensive"))
    model_start = next(i for i, l in enumerate(lines) if l.startswith("Valuation model"))
    block = lines[start:end] + [""] + lines[model_start:model_start + 5]
    return {
        "language": "r",
        "label": "R · tidyverse",
        "code": R_CODE,
        "output": "> source(\"R/02_analysis.R\")\n" + "\n".join(l.rstrip() for l in block).strip("\n"),
        "source": "Console output of Rscript R/02_analysis.R on the DVF sales 2021-2025",
    }


def main():
    outputs = {
        "captured_on": date.today().isoformat(),
        "sql": capture_sql(),
        "excel": capture_excel(),
        "powerbi": capture_powerbi(),
        "python": capture_python(),
        "r": capture_r(),
    }
    target = ROOT / "assets" / "data" / "code_outputs.js"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("// Generated by tools/capture_code_outputs.py - real outputs, do not edit by hand.\n"
                      f"window.CODE_OUTPUTS = {json.dumps(outputs, indent=2, ensure_ascii=False)};\n",
                      encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
    for key in ("sql", "excel", "powerbi", "python", "r"):
        print(f"===== {outputs[key]['label']}")
        print(outputs[key]["output"])
        print()


if __name__ == "__main__":
    main()
